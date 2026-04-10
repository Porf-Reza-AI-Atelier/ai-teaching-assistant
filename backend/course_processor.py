import os
import re
import time
import uuid
from pathlib import Path
from typing import List, Dict, Tuple
from dataclasses import dataclass
from llama_index.core import Document, VectorStoreIndex, Settings
from llama_index.core.embeddings import resolve_embed_model
from llama_index.core.text_splitter import SentenceSplitter
from llama_index.vector_stores.qdrant import QdrantVectorStore
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams, PointStruct
import PyPDF2
import tiktoken
from datetime import datetime
from docx import Document as DocxDocument
from pptx import Presentation

# Optional imports - gracefully handle if not available
try:
    import mammoth  # Better DOCX text extraction
    MAMMOTH_AVAILABLE = True
except ImportError:
    MAMMOTH_AVAILABLE = False
    print("ℹmammoth not available - using python-docx only for DOCX processing")

try:
    import magic  # File type detection
    MAGIC_AVAILABLE = True
except ImportError:
    MAGIC_AVAILABLE = False
    print("python-magic not available - using file extensions for type detection")

@dataclass
class CourseStructure:
    course_id: str
    course_name: str
    lessons: List[Dict]  # [{"lesson_id": "lesson_1", "lesson_name": "Variables", "documents": [...]}]
    general_documents: List[Dict]  # Documents not in specific lessons (syllabus, etc.)

class EnhancedCourseProcessor:
    def __init__(self):
        print("🔧 Initializing Enhanced Course Processor...")
        
        # Use local embeddings for cost efficiency
        Settings.embed_model = resolve_embed_model("local:BAAI/bge-small-en-v1.5")
        from llama_index.llms.groq import Groq
        Settings.llm = Groq(model="llama3-70b-8192", api_key=os.getenv("GROQ_API_KEY"))
        
        # Initialize enhanced text splitter for contextual chunking
        self.text_splitter = SentenceSplitter(
            chunk_size=512,
            chunk_overlap=50,
            paragraph_separator="\n\n\n",
            secondary_chunking_regex=r"[.!?]+\s+",
            tokenizer=tiktoken.get_encoding("cl100k_base").encode
        )
        
        self.qdrant_client = QdrantClient(
            url=os.getenv("QDRANT_URL"),
            api_key=os.getenv("QDRANT_API_KEY")
        )
        
        try:
            collections = self.qdrant_client.get_collections()
            print(f"✅ Qdrant connected. Existing collections: {[col.name for col in collections.collections]}")
        except Exception as e:
            print(f"❌ Qdrant connection failed: {e}")
            raise e

    def parse_course_structure(self, base_path: str) -> List[CourseStructure]:
        """
        Parse the downloaded Google Drive structure into course objects
        Expected structure: base_path/Course Name/Lesson X - Topic/documents.pdf
        Also handles general files outside lesson folders
        """
        print(f"📁 Parsing course structure from: {base_path}")
        courses = []
        
        base_dir = Path(base_path)
        if not base_dir.exists():
            raise ValueError(f"Path does not exist: {base_path}")
        
        # Iterate through course directories
        for course_dir in base_dir.iterdir():
            if not course_dir.is_dir():
                continue
                
            course_name = course_dir.name
            course_id = self._generate_course_id(course_name)
            lessons = []
            general_documents = []
            
            print(f"📚 Processing course: {course_name} (ID: {course_id})")
            
            # First, collect all items in course directory
            all_items = list(course_dir.iterdir())
            lesson_dirs = []
            
            # Separate lesson directories from other content
            for item in all_items:
                if item.is_dir() and self._is_lesson_directory(item.name):
                    lesson_dirs.append(item)
                elif item.is_file() and self._is_supported_file(item):
                    # General course document (syllabus, overview, etc.)
                    general_documents.append({
                        "file_path": str(item),
                        "file_name": item.name,
                        "file_type": item.suffix.lower(),
                        "category": "general"
                    })
                elif item.is_dir():
                    # Non-lesson directory, check for documents inside
                    folder_docs = self._extract_documents_from_folder(item, "general")
                    general_documents.extend(folder_docs)
            
            # Process lesson directories
            lesson_dirs = sorted(lesson_dirs, key=lambda x: self._extract_lesson_order(x.name))
            for lesson_dir in lesson_dirs:
                lesson_info = self._parse_lesson_info(lesson_dir.name)
                documents = self._extract_documents_from_folder(lesson_dir, "lesson")
                
                if documents:  # Only add lessons with documents
                    lessons.append({
                        "lesson_id": lesson_info["lesson_id"],
                        "lesson_name": lesson_info["lesson_name"],
                        "lesson_order": lesson_info["lesson_order"],
                        "documents": documents
                    })
                    print(f"  📖 Lesson {lesson_info['lesson_order']}: {lesson_info['lesson_name']} ({len(documents)} docs)")
            
            # Report general documents
            if general_documents:
                print(f"  📄 General documents: {len(general_documents)} files")
                for doc in general_documents:
                    print(f"    • {doc['file_name']}")
            
            if lessons or general_documents:  # Add course if it has any content
                courses.append(CourseStructure(
                    course_id=course_id,
                    course_name=course_name,
                    lessons=lessons,
                    general_documents=general_documents
                ))
        
        total_lessons = sum(len(c.lessons) for c in courses)
        total_general = sum(len(c.general_documents) for c in courses)
        print(f"✅ Parsed {len(courses)} courses with {total_lessons} lessons and {total_general} general documents")
        return courses

    def _is_lesson_directory(self, dir_name: str) -> bool:
        """Check if directory name represents a lesson"""
        patterns = [
            r'^(?:Lesson|Week|Chapter|Unit|Module)\s*\d+',
            r'^\d+\s*[-–—]',
            r'^(?:Lesson|Week|Chapter|Unit|Module)\s*\d+\s*[-–—:]'
        ]
        
        for pattern in patterns:
            if re.match(pattern, dir_name, re.IGNORECASE):
                return True
        return False

    def _is_supported_file(self, file_path: Path) -> bool:
        """Check if file type is supported"""
        supported_extensions = {'.pdf', '.txt', '.md', '.docx', '.pptx', '.doc', '.ppt'}
        return file_path.suffix.lower() in supported_extensions

    def _extract_lesson_order(self, dir_name: str) -> int:
        """Extract lesson order number for sorting"""
        patterns = [
            r'^(?:Lesson|Week|Chapter|Unit|Module)\s*(\d+)',
            r'^(\d+)\s*[-–—]'
        ]
        
        for pattern in patterns:
            match = re.match(pattern, dir_name, re.IGNORECASE)
            if match:
                return int(match.group(1))
        
        return 999  # Put unordered items at end

    def _extract_documents_from_folder(self, folder_path: Path, category: str) -> List[Dict]:
        """Extract all supported documents from a folder"""
        documents = []
        
        # Get all files (including in subdirectories for non-lesson folders)
        if category == "lesson":
            # For lesson folders, only look at direct files
            files = [f for f in folder_path.iterdir() if f.is_file()]
        else:
            # For general folders, look recursively
            files = [f for f in folder_path.rglob("*") if f.is_file()]
        
        for file_path in files:
            if self._is_supported_file(file_path):
                documents.append({
                    "file_path": str(file_path),
                    "file_name": file_path.name,
                    "file_type": file_path.suffix.lower(),
                    "category": category,
                    "subfolder": file_path.parent.name if file_path.parent != folder_path else None
                })
        
        return documents

    def _generate_course_id(self, course_name: str) -> str:
        """Generate a clean course ID from course name"""
        # Extract course code if present (e.g., "CS101 - Intro Programming" -> "cs101")
        course_code_match = re.match(r'^([A-Z]{2,4}\d{3})', course_name, re.IGNORECASE)
        if course_code_match:
            return course_code_match.group(1).lower()
        
        # Otherwise, create ID from first words
        words = re.findall(r'\w+', course_name.lower())
        return '_'.join(words[:3])  # Take first 3 words

    def _parse_lesson_info(self, lesson_dir_name: str) -> Dict:
        """Parse lesson directory name to extract lesson info"""
        # Expected format: "Lesson 1 - Variables" or "1 - Variables" or "Week 1 - Introduction"
        
        # Try different patterns
        patterns = [
            r'^(?:Lesson|Week|Chapter|Unit)\s*(\d+)\s*[-–—]\s*(.+)$',
            r'^(\d+)\s*[-–—]\s*(.+)$',
            r'^(?:Lesson|Week|Chapter|Unit)\s*(\d+)(?:\s*[:-]\s*(.+))?$'
        ]
        
        for pattern in patterns:
            match = re.match(pattern, lesson_dir_name, re.IGNORECASE)
            if match:
                lesson_num = int(match.group(1))
                lesson_name = match.group(2).strip() if len(match.groups()) > 1 and match.group(2) else f"Lesson {lesson_num}"
                return {
                    "lesson_id": f"lesson_{lesson_num}",
                    "lesson_name": lesson_name,
                    "lesson_order": lesson_num
                }
        
        # Fallback: use directory name as-is
        return {
            "lesson_id": re.sub(r'[^\w]', '_', lesson_dir_name.lower()),
            "lesson_name": lesson_dir_name,
            "lesson_order": 999  # Put unordered lessons at end
        }

    def extract_text_from_document(self, file_path: str) -> str:
        """Extract text from various document formats"""
        file_path = Path(file_path)
        
        try:
            if file_path.suffix.lower() == '.pdf':
                return self._extract_text_from_pdf(str(file_path))
            elif file_path.suffix.lower() in ['.txt', '.md']:
                return self._extract_text_from_text(str(file_path))
            elif file_path.suffix.lower() in ['.docx', '.doc']:
                return self._extract_text_from_docx(str(file_path))
            elif file_path.suffix.lower() in ['.pptx', '.ppt']:
                return self._extract_text_from_pptx(str(file_path))
            else:
                print(f"⚠️ Unsupported file type: {file_path.suffix}")
                return ""
        except Exception as e:
            print(f"❌ Error extracting text from {file_path}: {e}")
            return ""

    def _extract_text_from_text(self, file_path: str) -> str:
        """Extract text from TXT/MD files"""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            print(f"⚠️ Warning: Could not read text file {file_path}: {e}")
            return ""

    def _extract_text_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX files with multiple methods"""
        text = ""
        
        try:
            # Method 1: Try mammoth for better formatting preservation (if available)
            if MAMMOTH_AVAILABLE:
                try:
                    with open(file_path, "rb") as docx_file:
                        result = mammoth.extract_raw_text(docx_file)
                        text = result.value
                        if result.messages:
                            print(f"ℹ️ Mammoth extraction warnings for {file_path}: {len(result.messages)} messages")
                    print(f"✅ Extracted text from {file_path} using mammoth")
                    return text
                except Exception as mammoth_error:
                    print(f"⚠️ Mammoth extraction failed for {file_path}: {mammoth_error}")
            else:
                print(f"Mammoth not available, using python-docx for {file_path}")
                
            # Method 2: Use python-docx (always available as it's required)
            try:
                doc = DocxDocument(file_path)
                paragraphs = []
                
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        paragraphs.append(paragraph.text)
                
                # Also extract text from tables
                for table in doc.tables:
                    for row in table.rows:
                        row_text = []
                        for cell in row.cells:
                            if cell.text.strip():
                                row_text.append(cell.text.strip())
                        if row_text:
                            paragraphs.append(" | ".join(row_text))
                
                text = "\n\n".join(paragraphs)
                print(f"✅ Extracted text from {file_path} using python-docx")
                return text
                
            except Exception as docx_error:
                print(f"❌ python-docx extraction failed for {file_path}: {docx_error}")
                return ""
                
        except Exception as e:
            print(f"❌ Error extracting DOCX {file_path}: {e}")
            return ""

    def _extract_text_from_pptx(self, file_path: str) -> str:
        """Extract text from PPTX files"""
        try:
            presentation = Presentation(file_path)
            text_content = []
            
            for slide_num, slide in enumerate(presentation.slides, 1):
                slide_text = [f"\n--- Slide {slide_num} ---"]
                
                # Extract text from shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                    
                    # Extract text from tables in slides
                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            row_text = []
                            for cell in row.cells:
                                if cell.text.strip():
                                    row_text.append(cell.text.strip())
                            if row_text:
                                slide_text.append(" | ".join(row_text))
                
                if len(slide_text) > 1:  # More than just the slide header
                    text_content.extend(slide_text)
            
            final_text = "\n".join(text_content)
            print(f"✅ Extracted text from {file_path} ({len(presentation.slides)} slides)")
            return final_text
            
        except Exception as e:
            print(f"❌ Error extracting PPTX {file_path}: {e}")
            return ""

    def _extract_text_from_pdf(self, pdf_path: str) -> str:
        """Extract text from PDF with better error handling"""
        try:
            with open(pdf_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                text = ""
                for page_num, page in enumerate(reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text.strip():  # Only add non-empty pages
                            text += f"\n--- Page {page_num + 1} ---\n"
                            text += page_text + "\n"
                    except Exception as e:
                        print(f"⚠️ Warning: Could not extract page {page_num + 1} from {pdf_path}: {e}")
                        continue
                return text
        except Exception as e:
            print(f"❌ Error extracting text from {pdf_path}: {e}")
            return ""

    def create_contextual_chunks(self, text: str, metadata: Dict) -> List[Document]:
        """Create contextually-aware document chunks"""
        if not text.strip():
            return []
        
        # Use enhanced text splitter for better context preservation
        chunks = self.text_splitter.split_text(text)
        
        documents = []
        for i, chunk in enumerate(chunks):
            if not chunk.strip():
                continue
                
            # Enhanced metadata for each chunk
            chunk_metadata = {
                **metadata,
                "chunk_index": i,
                "total_chunks": len(chunks),
                "chunk_type": "contextual",
                "text_length": len(chunk),
                "processing_timestamp": datetime.now().isoformat()
            }
            
            # Add contextual information for better retrieval
            enhanced_text = self._add_contextual_prefix(chunk, metadata)
            
            documents.append(Document(
                text=enhanced_text,
                metadata=chunk_metadata
            ))
        
        return documents

    def _add_contextual_prefix(self, chunk_text: str, metadata: Dict) -> str:
        """Add contextual information to improve retrieval"""
        if metadata.get('lesson_order', 0) == 0:
            # General document
            context_prefix = f"""Course: {metadata['course_name']}
General Materials: {metadata['document_name']}
{f"Folder: {metadata['subfolder']}" if metadata.get('subfolder') else ""}

Content:
"""
        else:
            # Lesson document
            context_prefix = f"""Course: {metadata['course_name']}
Lesson {metadata['lesson_order']}: {metadata['lesson_name']}
Document: {metadata['document_name']}

Content:
"""
        return context_prefix + chunk_text

    def process_course_structure(self, courses: List[CourseStructure], force_recreate: bool = False):
        """Process entire course structure and store in vector database"""
        print(f"🚀 Processing {len(courses)} courses...")
        
        for course in courses:
            print(f"\n📚 Processing course: {course.course_name}")
            collection_name = f"course_{course.course_id}"
            
            # Handle existing collections
            if force_recreate:
                self._recreate_collection(collection_name)
            else:
                self._ensure_collection_exists(collection_name)
            
            # Process all documents in course
            all_documents = []
            total_lesson_docs = sum(len(lesson["documents"]) for lesson in course.lessons)
            total_general_docs = len(course.general_documents)
            total_docs = total_lesson_docs + total_general_docs
            processed_docs = 0
            
            print(f"  📊 Total documents to process: {total_docs} ({total_lesson_docs} lesson + {total_general_docs} general)")
            
            # Process general documents first
            if course.general_documents:
                print(f"  📄 Processing {len(course.general_documents)} general documents...")
                for doc_info in course.general_documents:
                    try:
                        # Extract text
                        text = self.extract_text_from_document(doc_info["file_path"])
                        if not text.strip():
                            print(f"    ⚠️ Skipping empty document: {doc_info['file_name']}")
                            continue
                        
                        # Create metadata for general documents
                        metadata = {
                            "course_id": course.course_id,
                            "course_name": course.course_name,
                            "lesson_id": "general",
                            "lesson_name": "General Course Materials",
                            "lesson_order": 0,  # General docs come first
                            "document_name": doc_info["file_name"],
                            "document_path": doc_info["file_path"],
                            "document_type": doc_info["file_type"],
                            "category": doc_info.get("category", "general"),
                            "subfolder": doc_info.get("subfolder")
                        }
                        
                        # Create contextual chunks
                        documents = self.create_contextual_chunks(text, metadata)
                        all_documents.extend(documents)
                        
                        processed_docs += 1
                        print(f"    ✅ {doc_info['file_name']}: {len(documents)} chunks")
                        
                    except Exception as e:
                        print(f"    ❌ Error processing {doc_info['file_name']}: {e}")
                        continue
            
            # Process lesson documents
            for lesson in course.lessons:
                print(f"  📖 Processing lesson: {lesson['lesson_name']}")
                
                for doc_info in lesson["documents"]:
                    try:
                        # Extract text
                        text = self.extract_text_from_document(doc_info["file_path"])
                        if not text.strip():
                            print(f"    ⚠️ Skipping empty document: {doc_info['file_name']}")
                            continue
                        
                        # Create metadata for lesson documents
                        metadata = {
                            "course_id": course.course_id,
                            "course_name": course.course_name,
                            "lesson_id": lesson["lesson_id"],
                            "lesson_name": lesson["lesson_name"],
                            "lesson_order": lesson["lesson_order"],
                            "document_name": doc_info["file_name"],
                            "document_path": doc_info["file_path"],
                            "document_type": doc_info["file_type"],
                            "category": doc_info.get("category", "lesson"),
                            "subfolder": doc_info.get("subfolder")
                        }
                        
                        # Create contextual chunks
                        documents = self.create_contextual_chunks(text, metadata)
                        all_documents.extend(documents)
                        
                        processed_docs += 1
                        print(f"    ✅ {doc_info['file_name']}: {len(documents)} chunks")
                        
                    except Exception as e:
                        print(f"    ❌ Error processing {doc_info['file_name']}: {e}")
                        continue
            
            # Store all documents for this course
            if all_documents:
                self._store_documents_in_collection(collection_name, all_documents)
                print(f"✅ Course {course.course_name}: {len(all_documents)} total chunks stored ({processed_docs}/{total_docs} docs processed)")
            else:
                print(f"⚠️ No documents processed for course {course.course_name}")

    def _recreate_collection(self, collection_name: str):
        """Delete and recreate collection"""
        try:
            collections = self.qdrant_client.get_collections()
            existing_collections = [col.name for col in collections.collections]
            
            if collection_name in existing_collections:
                print(f"🗑️ Deleting existing collection: {collection_name}")
                self.qdrant_client.delete_collection(collection_name)
            
            print(f"🆕 Creating collection: {collection_name}")
            self.qdrant_client.create_collection(
                collection_name=collection_name,
                vectors_config=VectorParams(size=384, distance=Distance.COSINE)
            )
        except Exception as e:
            print(f"❌ Error recreating collection {collection_name}: {e}")
            raise

    def _ensure_collection_exists(self, collection_name: str):
        """Ensure collection exists, create if not"""
        try:
            collections = self.qdrant_client.get_collections()
            existing_collections = [col.name for col in collections.collections]
            
            if collection_name not in existing_collections:
                print(f"🆕 Creating new collection: {collection_name}")
                self.qdrant_client.create_collection(
                    collection_name=collection_name,
                    vectors_config=VectorParams(size=384, distance=Distance.COSINE)
                )
            else:
                print(f"📁 Using existing collection: {collection_name}")
        except Exception as e:
            print(f"❌ Error ensuring collection {collection_name}: {e}")
            raise

    def _store_documents_in_collection(self, collection_name: str, documents: List[Document]):
        """Store documents in Qdrant collection with improved error handling and proper point IDs"""
        try:
            print(f"📤 Storing {len(documents)} documents in {collection_name}...")
            
            # Try LlamaIndex method first (more robust)
            try:
                vector_store = QdrantVectorStore(
                    client=self.qdrant_client,
                    collection_name=collection_name
                )
                
                print(f"🔄 Attempting LlamaIndex storage...")
                index = VectorStoreIndex.from_documents(
                    documents,
                    vector_store=vector_store,
                    show_progress=True
                )
                
                # Wait for async operations and verify multiple times
                print("⏳ Waiting for storage to complete...")
                for attempt in range(10):  # Try more attempts
                    time.sleep(2)  # Wait longer between checks
                    count_result = self.qdrant_client.count(collection_name)
                    print(f"📊 Verification attempt {attempt + 1}: {count_result.count} vectors")
                    
                    if count_result.count > 0:
                        print(f"✅ LlamaIndex storage successful: {count_result.count} vectors stored!")
                        return
                
                print(f"⚠️ LlamaIndex storage verification failed after 10 attempts")
                
            except Exception as e:
                print(f"❌ LlamaIndex storage failed: {e}")
            
            # If LlamaIndex fails, use direct storage with proper point IDs
            print(f"🔄 Using direct Qdrant storage method...")
            self._direct_vector_storage(collection_name, documents)
            
        except Exception as e:
            print(f"❌ All storage methods failed: {e}")
            raise

    def _direct_vector_storage(self, collection_name: str, documents: List[Document]):
        """Direct vector storage with proper Qdrant point format and IDs"""
        try:
            print(f"🔧 Direct storage: Processing {len(documents)} documents...")
            
            embed_model = Settings.embed_model
            points = []
            
            # Generate embeddings for all documents first
            for i, doc in enumerate(documents):
                try:
                    # Generate embedding
                    embedding = embed_model.get_text_embedding(doc.text)
                    
                    # Create enhanced payload with text content for LlamaIndex compatibility
                    enhanced_payload = {
                        **doc.metadata,
                        "_node_content": doc.text,  # Store text content for retrieval
                        "_node_type": "TextNode",   # Indicate node type
                        "text": doc.text            # Also store as 'text' field
                    }
                    
                    # Create proper Qdrant point with integer ID
                    point = PointStruct(
                        id=i,  # Use simple integer ID (valid for Qdrant)
                        vector=embedding,
                        payload=enhanced_payload
                    )
                    points.append(point)
                    
                    if (i + 1) % 50 == 0:
                        print(f"📊 Generated embeddings for {i + 1}/{len(documents)} documents")
                        
                except Exception as e:
                    print(f"⚠️ Error processing document {i}: {e}")
                    continue
            
            if not points:
                print(f"❌ No valid points to store")
                return
            
            print(f"📦 Storing {len(points)} points in batches...")
            
            # Store in smaller batches with error handling
            batch_size = 50  # Smaller batches for reliability
            total_stored = 0
            
            for i in range(0, len(points), batch_size):
                batch = points[i:i + batch_size]
                batch_num = i // batch_size + 1
                total_batches = (len(points) + batch_size - 1) // batch_size
                
                try:
                    # Use upsert with proper error handling
                    self.qdrant_client.upsert(
                        collection_name=collection_name,
                        points=batch
                    )
                    total_stored += len(batch)
                    print(f"📤 Stored batch {batch_num}/{total_batches} ({len(batch)} vectors) - Total: {total_stored}")
                    time.sleep(1)  # Small delay between batches
                    
                except Exception as e:
                    print(f"❌ Error storing batch {batch_num}: {e}")
                    # Try individual points in failed batch
                    print(f"🔄 Retrying batch {batch_num} point by point...")
                    for j, point in enumerate(batch):
                        try:
                            self.qdrant_client.upsert(
                                collection_name=collection_name,
                                points=[point]
                            )
                            total_stored += 1
                        except Exception as point_error:
                            print(f"❌ Failed to store point {i + j}: {point_error}")
                            continue
            
            # Final verification with multiple attempts
            print(f"🔍 Verifying storage...")
            for attempt in range(5):
                time.sleep(2)
                count_result = self.qdrant_client.count(collection_name)
                print(f"📊 Verification attempt {attempt + 1}: {count_result.count} vectors")
                
                if count_result.count > 0:
                    print(f"✅ Direct storage successful: {count_result.count} vectors stored!")
                    return
            
            print(f"❌ Direct storage verification failed - no vectors found in collection")
            
        except Exception as e:
            print(f"❌ Direct storage failed: {e}")
            import traceback
            traceback.print_exc()
            raise

# Usage Example
if __name__ == "__main__":
    # Initialize processor
    processor = EnhancedCourseProcessor()
    
    # Parse course structure from downloaded Google Drive folder
    courses = processor.parse_course_structure("./downloaded_courses")
    
    # Process and store all courses
    processor.process_course_structure(courses, force_recreate=True)